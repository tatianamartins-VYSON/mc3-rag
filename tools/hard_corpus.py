"""Gera um corpus MAIOR e mais difícil que a amostra oficial, com perguntas e citações esperadas.

Objetivo: testar, sem GPU, o que quebra quando o corpus cresce — leitura dos arquivos,
armadilhas (vazio, sem permissão, criptografado, tipo desconhecido), revisões superadas,
distratores quase iguais, cadeias de dois arquivos (multi-hop) e perguntas sem resposta.

Uso:
    python3 tools/hard_corpus.py /tmp/hard   # cria /tmp/hard/corpus e /tmp/hard/questions.json
"""
import csv
import json
import os
import random
import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

random.seed(27)
OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/hard")
C = OUT / "corpus"
Q: list[dict] = []


def pdf(rel, pages):
    p = C / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(p), pagesize=A4)
    for lines in pages:
        y = 800
        for ln in lines:
            c.drawString(40, y, ln)
            y -= 16
        c.showPage()
    c.save()


def txt(rel, body):
    p = C / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(body, encoding="utf-8")


def q(query, answer, cites, cat):
    Q.append({"n": len(Q) + 1, "query": query, "expected_answer": answer, "expected_citations": cites, "category": cat})


# 1) datasheets: 12 produtos, cada um com revisões; só a última vale
for i, prod in enumerate(["TQ-10", "TQ-20", "TQ-30", "TQ-40", "TQ-50", "TQ-60", "TQ-70", "TQ-80", "TQ-90", "VX-12", "VX-24", "VX-48"]):
    tj_old, tj = 85 + i * 2, 90 + i * 2
    slug = prod.lower().replace("-", "")
    pdf(f"specs/{slug}_datasheet_r1_SUPERSEDED.pdf",
        [[f"{prod} Datasheet revision R1 - SUPERSEDED by R2, do not use", f"Maximum junction temperature: {tj_old} C",
          f"Supply voltage: {1.0 + i/10:.1f} V", "Package: BGA-484"]])
    pdf(f"specs/{slug}_datasheet_r2.pdf",
        [[f"{prod} Datasheet revision R2 (current)", "Section 1 - Overview", f"The {prod} is a compute module."],
         [f"{prod} Datasheet R2 - Section 4 Thermal", f"Maximum junction temperature: {tj} C", f"Thermal resistance: {0.30 + i/100:.2f} C/W",
          f"Supply voltage: {1.1 + i/10:.1f} V"]])
    if prod in ("TQ-40", "VX-24", "TQ-80"):
        q(f"What is the maximum junction temperature of the {prod}?", str(tj), [f"specs/{slug}_datasheet_r2.pdf"], "pdf multi-page + superseded trap + 11 distractors")

# 2) roadmap em Word com tabela
d = Document()
d.add_heading("Roadmap FY27", 1)
d.add_paragraph("Customer sampling schedule by product. Dates may move; this table is authoritative.")
t = d.add_table(rows=1, cols=3)
t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Product", "Sampling", "Volume"
for prod, s, v in [("TQ-60", "Q3 FY27", "Q1 FY28"), ("TQ-70", "Q4 FY27", "Q2 FY28"), ("VX-48", "Q2 FY27", "Q4 FY27"), ("TQ-90", "Q1 FY28", "Q3 FY28")]:
    r = t.add_row().cells
    r[0].text, r[1].text, r[2].text = prod, s, v
(C / "planning").mkdir(parents=True, exist_ok=True)
d.save(C / "planning/roadmap_fy27.docx")
d2 = Document()
d2.add_paragraph("DRAFT roadmap FY26 (obsolete). TQ-60 sampling was planned for Q1 FY27 at that time.")
d2.save(C / "planning/roadmap_fy26_draft_OBSOLETE.docx")
q("In which quarter does the TQ-70 enter customer sampling?", "Q4 FY27", ["planning/roadmap_fy27.docx"], "docx table + obsolete draft trap")
q("When does the VX-48 reach volume production?", "Q4 FY27", ["planning/roadmap_fy27.docx"], "docx table, second column")

# 3) planilha com várias abas (peças, custos, prazos)
wb = Workbook()
ws = wb.active
ws.title = "parts"
ws.append(["part_no", "description", "product"])
parts = {}
for k in range(400):
    pn = f"PN-{70000 + k}"
    parts[pn] = random.choice(["fan", "heatsink", "bracket", "cable", "PSU", "riser"])
    ws.append([pn, parts[pn], random.choice(["TQ-40", "TQ-60", "VX-24"])])
ws2 = wb.create_sheet("costs")
ws2.append(["part_no", "unit_cost_usd"])
for k in range(400):
    ws2.append([f"PN-{70000 + k}", round(5 + k * 0.37, 2)])
ws3 = wb.create_sheet("lead_times")
ws3.append(["part_no", "lead_time_weeks"])
for k in range(400):
    ws3.append([f"PN-{70000 + k}", 2 + k % 17])
(C / "support").mkdir(parents=True, exist_ok=True)
wb.save(C / "support/rma_parts.xlsx")
q("What is the unit cost in USD of part PN-70321?", f"{round(5 + 321 * 0.37, 2)}", ["support/rma_parts.xlsx"], "xlsx second sheet among 1200 rows")
q("What is the lead time in weeks for part PN-70388?", str(2 + 388 % 17), ["support/rma_parts.xlsx"], "xlsx third sheet")

# 4) banco de bugs CSV grande + log que referencia o ticket (cadeia de 2 arquivos)
with open(C / "support/bug_database.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["ticket", "component", "fix_version", "status"])
    for k in range(3000):
        w.writerow([f"ORR-{1000 + k}", random.choice(["ingest", "scheduler", "firmware", "ui"]), f"4.{k % 9}.{k % 13}", random.choice(["open", "fixed"])])
lines = [f"2026-09-02T0{h % 10}:{m:02d}:00 INFO request ok latency_ms={random.randint(5, 90)}" for h in range(6) for m in range(60)]
lines.insert(211, "2026-09-02T03:31:00 ERROR worker crashed, tracked as ORR-2847 (see bug database)")
(C / "logs").mkdir(parents=True, exist_ok=True)
txt("logs/prod_inference_2026-09-02.log", "\n".join(lines))
txt("logs/prod_inference_2026-09-01.log", "\n".join(l.replace("09-02", "09-01") for l in lines if "ERROR" not in l))
q("The production log from 2026-09-02 shows a worker crash. In which version is the related ticket fixed?",
  f"4.{1847 % 9}.{1847 % 13}", ["logs/prod_inference_2026-09-02.log", "support/bug_database.csv"], "multi-hop log -> csv (3000 rows)")

# 5) código-fonte com constantes
txt("engineering/ingest_service.py", "\n".join([
    '"""Ingest service."""', "DEFAULT_BATCH_SIZE = 256", "MAX_RETRIES = 7", "TIMEOUT_SECONDS = 45",
    "def run(batch_size=DEFAULT_BATCH_SIZE):", "    return batch_size"]))
txt("engineering/legacy/ingest_service_v1.py", "DEFAULT_BATCH_SIZE = 64  # legacy, removed in 4.0\nMAX_RETRIES = 3\n")
q("How many retries does the ingest service make by default?", "7", ["engineering/ingest_service.py"], "py constant + legacy distractor")

# 6) release notes
txt("engineering/meridian_release_notes.txt", "\n".join(
    [f"Meridian {v} - {note}" for v, note in [("4.3.0", "new scheduler"), ("4.3.1", "fixes ORR-1201"), ("4.3.2", "raises inference timeout to 120 s"), ("4.4.0", "drops Python 3.9")]]))
q("Which Meridian release raised the inference timeout?", "4.3.2", ["engineering/meridian_release_notes.txt"], "txt release notes")

# 7) muitos documentos de ruído (o corpus oficial é maior)
for k in range(120):
    txt(f"archive/notes/meeting_{k:03d}.txt", f"Meeting {k}: discussed TQ-40 thermals, roadmap and costs. No decisions recorded. Next sync in two weeks.")

# 7b) imagens: a resposta só existe como texto dentro da figura (precisa do VLM)
from PIL import Image, ImageDraw, ImageFont
try:
    fnt = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 34)
except OSError:
    fnt = ImageFont.load_default()
img = Image.new("RGB", (900, 500), "white")
dr = ImageDraw.Draw(img)
dr.rectangle([40, 40, 860, 460], outline="black", width=4)
for k, (pin, sig) in enumerate([("P1", "GND"), ("P2", "12V"), ("P3", "PWM_FAN"), ("P4", "TACH")]):
    dr.text((80, 80 + k * 80), f"VX-24 backplane pin {pin}: {sig}", fill="black", font=fnt)
(C / "specs").mkdir(parents=True, exist_ok=True)
img.save(C / "specs/vx24_backplane_pinout.png")
ph = Image.new("RGB", (800, 400), (200, 200, 190))
dp = ImageDraw.Draw(ph)
dp.rectangle([120, 120, 680, 280], fill="white", outline="black", width=3)
dp.text((150, 150), "ASSET TAG: VX-AT-55821", fill="black", font=fnt)
dp.text((150, 210), "Rack R14 / Slot 7", fill="black", font=fnt)
ph.save(C / "support/rack_asset_label.jpg", quality=85)
q("Which signal is on pin P3 of the VX-24 backplane?", "PWM_FAN", ["specs/vx24_backplane_pinout.png"], "png: answer only inside the image (VLM)")
q("What asset tag is printed on the label in rack R14?", "VX-AT-55821", ["support/rack_asset_label.jpg"], "jpg photo: answer only inside the image (VLM)")

# 8) armadilhas
(C / "archive/empty").mkdir(parents=True, exist_ok=True)
txt("support/unknown.xyzbin", "\x00\x01binary-ish")
w = PdfWriter()
w.add_blank_page(200, 200)
w.encrypt("segredo")
with open(C / "specs/confidential_ENCRYPTED.pdf", "wb") as f:
    w.write(f)
txt("support/locked.txt", "TQ-40 secret junction temperature 999 C")
try:
    os.chmod(C / "support/locked.txt", 0)
except OSError:
    pass

# 9) perguntas sem resposta (devem voltar vazias)
q("What is the maximum junction temperature of the TQ-99?", "", [], "negative: product does not exist")
q("Who is the CEO of the company?", "", [], "negative: not in corpus")
q("What is the unit cost of part PN-99999?", "", [], "negative: id not in spreadsheet")

OUT.mkdir(parents=True, exist_ok=True)
(OUT / "questions.json").write_text(json.dumps({"queries": Q}, ensure_ascii=False, indent=1), encoding="utf-8")
n = sum(1 for _ in C.rglob("*") if _.is_file())
print(f"corpus: {n} arquivos em {C}; {len(Q)} perguntas em {OUT/'questions.json'}")
