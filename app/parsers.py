"""Defensive corpus walk and per-type parsers for Mini-Challenge 3.

Every parser returns a list of Chunk. Any exception is caught per file and the
walk continues; unreadable, encrypted and unknown files are recorded in the
index manifest instead of aborting the pass.

Images are NOT parsed here: they are listed as pending and transcribed by the
vision model in server.py during --index (see index_images()).
"""

from __future__ import annotations

import csv
import io
import os
import re
import sys
import zipfile
from dataclasses import dataclass, asdict
from pathlib import Path

TEXT_EXT = {".txt", ".log", ".md", ".py", ".json", ".yaml", ".yml", ".ini", ".cfg", ".toml", ".sh"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}

MAX_CHUNK_CHARS = 900


@dataclass
class Chunk:
    cid: str            # unique id, e.g. "specs/tq40_datasheet_r2.pdf#p1c0"
    file: str           # corpus-relative path with forward slashes
    kind: str           # pdf | docx | xlsx | csv | text | code | image
    text: str
    meta: str = ""      # page / sheet / line range, for the prompt

    def to_dict(self) -> dict:
        return asdict(self)


def _log(msg: str) -> None:
    print(f"[parsers] {msg}", file=sys.stderr, flush=True)


def _split(text: str, max_chars: int = MAX_CHUNK_CHARS) -> list[str]:
    """Split on blank lines first, then hard-wrap oversized paragraphs."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    out: list[str] = []
    buf = ""
    for p in paras:
        if len(buf) + len(p) + 1 <= max_chars:
            buf = (buf + "\n" + p).strip()
            continue
        if buf:
            out.append(buf)
        buf = ""
        while len(p) > max_chars:
            cut = p.rfind("\n", 0, max_chars)
            if cut < max_chars // 2:
                cut = max_chars
            out.append(p[:cut].strip())
            p = p[cut:].strip()
        buf = p
    if buf:
        out.append(buf)
    return out


# ----------------------------------------------------------------------------
# per-type parsers
# ----------------------------------------------------------------------------

def parse_pdf(path: Path, rel: str) -> list[Chunk]:
    chunks: list[Chunk] = []
    try:
        import pdfplumber
    except ImportError:
        pdfplumber = None
    if pdfplumber is not None:
        with pdfplumber.open(str(path)) as pdf:
            for pno, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                # tables as pipe-separated rows keep column adjacency
                try:
                    for tbl in page.extract_tables() or []:
                        rows = [" | ".join(str(c or "").strip() for c in row) for row in tbl if row]
                        if rows:
                            text += "\n\n" + "\n".join(rows)
                except Exception:  # noqa: BLE001 - tables are best effort
                    pass
                for i, part in enumerate(_split(text)):
                    chunks.append(Chunk(f"{rel}#p{pno}c{i}", rel, "pdf", part, f"page {pno}"))
        return chunks
    from pypdf import PdfReader  # fallback

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise PermissionError("encrypted pdf")
    for pno, page in enumerate(reader.pages, start=1):
        for i, part in enumerate(_split(page.extract_text() or "")):
            chunks.append(Chunk(f"{rel}#p{pno}c{i}", rel, "pdf", part, f"page {pno}"))
    return chunks


def parse_docx(path: Path, rel: str) -> list[Chunk]:
    """python-docx when available; raw XML fallback keeps table text too."""
    text = ""
    try:
        import docx  # python-docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n\n".join(parts)
    except Exception:  # noqa: BLE001
        with zipfile.ZipFile(str(path)) as z:
            xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
        xml = re.sub(r"</w:p>", "\n\n", xml)
        xml = re.sub(r"</w:tc>", " | ", xml)
        text = re.sub(r"<[^>]+>", "", xml)
    return [Chunk(f"{rel}#c{i}", rel, "docx", part) for i, part in enumerate(_split(text))]


def parse_xlsx(path: Path, rel: str) -> list[Chunk]:
    import openpyxl

    chunks: list[Chunk] = []
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = []
        header: list[str] | None = None
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v).strip() for v in row]
            if not any(vals):
                continue
            if header is None:
                header = vals
                rows.append(" | ".join(vals))
                continue
            # "header=value" pairs make single-row chunks self-describing
            pairs = [f"{h}={v}" for h, v in zip(header, vals) if v]
            rows.append(" | ".join(pairs) if pairs else " | ".join(vals))
        # one chunk per row block keeps the sheet name attached to each value
        block: list[str] = []
        n = 0
        for r in rows:
            block.append(r)
            if sum(len(x) for x in block) > MAX_CHUNK_CHARS:
                chunks.append(Chunk(f"{rel}#{ws.title}c{n}", rel, "xlsx", "\n".join(block), f"sheet {ws.title}"))
                block, n = [rows[0]] if header else [], n + 1
        if block and (len(block) > 1 or not header):
            chunks.append(Chunk(f"{rel}#{ws.title}c{n}", rel, "xlsx", "\n".join(block), f"sheet {ws.title}"))
    return chunks


def parse_csv(path: Path, rel: str) -> list[Chunk]:
    raw = path.read_bytes().decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(raw))
    rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        return []
    header = rows[0]
    lines = [" | ".join(header)]
    for r in rows[1:]:
        lines.append(" | ".join(f"{h}={v.strip()}" for h, v in zip(header, r) if v.strip()))
    chunks: list[Chunk] = []
    block: list[str] = [lines[0]]
    n = 0
    for line in lines[1:]:
        block.append(line)
        if sum(len(x) for x in block) > MAX_CHUNK_CHARS:
            chunks.append(Chunk(f"{rel}#c{n}", rel, "csv", "\n".join(block)))
            block, n = [lines[0]], n + 1
    if len(block) > 1:
        chunks.append(Chunk(f"{rel}#c{n}", rel, "csv", "\n".join(block)))
    return chunks


def parse_text(path: Path, rel: str, kind: str) -> list[Chunk]:
    raw = path.read_bytes()
    if b"\x00" in raw[:4096]:
        raise ValueError("binary content in a text-typed file")
    text = raw.decode("utf-8", errors="ignore")
    if kind == "code":
        # keep code line-oriented: constants and defaults sit on single lines
        lines = text.splitlines()
        chunks: list[Chunk] = []
        for i in range(0, len(lines), 40):
            part = "\n".join(lines[i:i + 40]).strip()
            if part:
                chunks.append(Chunk(f"{rel}#L{i + 1}", rel, "code", part, f"lines {i + 1}-{i + 40}"))
        return chunks
    if kind == "log":
        lines = [ln for ln in text.splitlines() if ln.strip()]
        chunks = []
        for i in range(0, len(lines), 25):
            chunks.append(Chunk(f"{rel}#L{i + 1}", rel, "log", "\n".join(lines[i:i + 25]), f"lines {i + 1}-{i + 25}"))
        return chunks
    return [Chunk(f"{rel}#c{i}", rel, "text", part) for i, part in enumerate(_split(text))]


# ----------------------------------------------------------------------------
# walk
# ----------------------------------------------------------------------------

def walk(corpus: Path) -> tuple[list[Chunk], list[str], list[dict]]:
    """Return (chunks, pending_images, manifest).

    manifest lists every file seen with status: parsed | image-pending |
    skipped-unknown | unreadable | encrypted | error.
    """
    chunks: list[Chunk] = []
    images: list[str] = []
    manifest: list[dict] = []
    corpus = corpus.resolve()

    for root, dirs, files in os.walk(corpus, onerror=lambda e: _log(f"walk error ignored: {e}")):
        dirs.sort()
        for name in sorted(files):
            path = Path(root) / name
            rel = path.relative_to(corpus).as_posix()
            ext = path.suffix.lower()
            entry = {"file": rel, "status": "parsed", "chunks": 0}
            try:
                if ext in IMAGE_EXT:
                    # verify it opens; transcription happens in the server
                    from PIL import Image

                    with Image.open(path) as im:
                        im.verify()
                    images.append(rel)
                    entry["status"] = "image-pending"
                elif ext == ".pdf":
                    got = parse_pdf(path, rel)
                    if not got:
                        raise ValueError("no extractable text (scanned or encrypted)")
                    chunks += got
                    entry["chunks"] = len(got)
                elif ext == ".docx":
                    got = parse_docx(path, rel)
                    chunks += got
                    entry["chunks"] = len(got)
                elif ext in {".xlsx", ".xlsm"}:
                    got = parse_xlsx(path, rel)
                    chunks += got
                    entry["chunks"] = len(got)
                elif ext == ".csv":
                    got = parse_csv(path, rel)
                    chunks += got
                    entry["chunks"] = len(got)
                elif ext in TEXT_EXT:
                    kind = "code" if ext == ".py" else ("log" if ext == ".log" else "text")
                    got = parse_text(path, rel, kind)
                    chunks += got
                    entry["chunks"] = len(got)
                else:
                    entry["status"] = "skipped-unknown"
            except PermissionError as e:
                entry["status"] = "unreadable"
                entry["error"] = str(e)[:200]
            except Exception as e:  # noqa: BLE001 - never let one file kill the walk
                msg = f"{type(e).__name__}: {e}"[:200]
                entry["status"] = "encrypted" if "crypt" in msg.lower() or "password" in msg.lower() else "error"
                entry["error"] = msg
            if entry["status"] != "parsed":
                _log(f"{entry['status']}: {rel} {entry.get('error', '')}")
            manifest.append(entry)
    return chunks, images, manifest
